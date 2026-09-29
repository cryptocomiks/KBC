"""Bank statement analysis: the red flags an analyst looks for in a client's transactions.

The file (CSV or Excel export of any bank) is read in memory, analysed and discarded:
nothing is stored. Columns are recognised from their headers in English, French, German
and Italian; amounts in Swiss (1'234.50), French (1 234,50) and German (1.234,50) formats.

Detectors (each alert lists the transactions behind it):
* structuring: amounts just below the identification threshold, or split over a few days;
* cash: deposits / withdrawals in cash, and cash declared as "none" in the KYC profile;
* round amounts: a large share of round transfers;
* pass-through: money that leaves within days of arriving, for about the same amount;
* geography: counterparties in FATF black / grey list, EU tax list, offshore, broadly
  sanctioned or high Basel AML Index countries (from the IBAN or the country column);
* sanctions / watchlists: counterparty names screened against the enabled lists;
* crypto: transfers to or from crypto-asset platforms;
* spikes: months far above the usual activity;
* profile: annualised volume above the amount declared in the KYC questionnaire;
* funnel: many different senders, money leaving to very few beneficiaries.
"""

from __future__ import annotations

import csv
import io
import re
import statistics
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from typing import Any

from unidecode import unidecode

from app.connectors.util import nationality_iso
from app.models import Entity, EntityType
from app.risk.config import get_country_risk, get_jurisdictions, get_risk_config

MAX_BYTES = 5_000_000
MAX_ROWS = 20_000
MAX_SCREENED = 60
BROADLY_SANCTIONED = {"KP", "IR", "SY", "CU", "RU", "BY"}
# Declared annual amounts (KYC questionnaire, in EUR): upper bound of each band
VOLUME_BANDS = {
    "lt_150k": (0, 150_000),
    "150k_1m": (150_000, 1_000_000),
    "1m_10m": (1_000_000, 10_000_000),
    "gt_10m": (10_000_000, None),
}

HEADERS: dict[str, tuple[str, ...]] = {
    "date": (
        "date",
        "booking date",
        "transaction date",
        "value date",
        "date comptable",
        "date operation",
        "date valeur",
        "datum",
        "buchungsdatum",
        "valutadatum",
        "data",
        "posting date",
    ),
    "amount": ("amount", "montant", "betrag", "importo", "value", "sum", "transaction amount"),
    "debit": (
        "debit",
        "withdrawal",
        "withdrawals",
        "paid out",
        "money out",
        "out",
        "soll",
        "belastung",
        "lastschrift",
        "retrait",
        "depense",
        "addebito",
        "uscite",
    ),
    "credit": (
        "credit",
        "deposit",
        "deposits",
        "paid in",
        "money in",
        "in",
        "haben",
        "gutschrift",
        "versement",
        "accredito",
        "entrate",
    ),
    "currency": ("currency", "devise", "ccy", "wahrung", "monnaie", "divisa", "valuta code"),
    "counterparty": (
        "counterparty",
        "counterparty name",
        "beneficiary",
        "payee",
        "payer",
        "name",
        "contrepartie",
        "beneficiaire",
        "donneur d'ordre",
        "emetteur",
        "auftraggeber",
        "empfanger",
        "empfaenger",
        "partner",
        "controparte",
        "ordering party",
    ),
    "iban": (
        "iban",
        "counterparty iban",
        "account",
        "counterparty account",
        "compte",
        "konto",
        "conto",
        "iban contrepartie",
    ),
    "country": ("country", "counterparty country", "pays", "land", "paese"),
    "description": (
        "description",
        "libelle",
        "details",
        "reference",
        "communication",
        "verwendungszweck",
        "motif",
        "text",
        "buchungstext",
        "causale",
        "narrative",
        "memo",
    ),
    "type": ("type", "transaction type", "nature", "operation", "art", "tipo"),
}
CASH = re.compile(
    r"\b(cash|atm|especes|espece|retrait (dab|gab)|versement especes|bargeld|bareinzahlung|barbezug|bancomat|contanti|prelievo|dab|gab)\b",
    re.I,
)
CRYPTO = re.compile(
    r"\b(binance|coinbase|kraken|bitstamp|bitpanda|bitfinex|okx|okex|kucoin|bybit|crypto\.com|gemini|huobi|htx|gate\.io|swissborg|bitvavo|coinmerce|nexo|blockchain\.com|paxful|localbitcoins|changelly|mexc|bitget)\b",
    re.I,
)
IBAN = re.compile(r"\b([A-Z]{2}\d{2}[A-Z0-9]{10,30})\b")
LEGAL_WORDS = re.compile(
    r"\b(ltd|limited|llc|inc|corp|gmbh|ag|sa|sarl|sas|srl|spa|bv|nv|plc|oy|ab|as|holding|holdings|trading|group|company|co|foundation|trust|bank|s\.a\.|s\.à r\.l\.)\b",
    re.I,
)


def _norm(h: Any) -> str:
    return re.sub(r"[^a-z' ]", " ", unidecode(str(h or "")).lower()).strip()


def _amount(v: Any) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, int | float):
        return float(v)
    s = str(v).strip()
    neg = s.startswith("(") and s.endswith(")") or s.endswith("-") or s.startswith("-")
    s = re.sub(r"[^\d,.'\s]", "", s).replace("'", "").replace(" ", "").replace(" ", "")
    if not s:
        return None
    if "," in s and "." in s:
        s = (
            s.replace(".", "").replace(",", ".")
            if s.rfind(",") > s.rfind(".")
            else s.replace(",", "")
        )
    elif "," in s:
        head, _, tail = s.rpartition(",")
        # "1,234" / "1,234,567": thousands separators; "12,50": decimal comma
        thousands = s.count(",") > 1 or (len(tail) == 3 and 0 < len(head) <= 3)
        s = s.replace(",", "") if thousands else s.replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")  # "1.234.567"
    try:
        val = float(s)
    except ValueError:
        return None
    return -val if neg else val


def _date(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()[:19]
    for fmt in (
        "%Y-%m-%d",
        "%d.%m.%Y",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y/%m/%d",
        "%d.%m.%y",
        "%d/%m/%y",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(s if " " in fmt or "T" in fmt else s[:10], fmt).date()
        except ValueError:
            continue
    return None


def _rows(data: bytes, filename: str) -> list[list[Any]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)][: MAX_ROWS + 20]
    text = data.decode("utf-8-sig", errors="strict") if _is_utf8(data) else data.decode("latin-1")
    sample = text[:5000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = max(",;\t|", key=sample.count)
    return list(csv.reader(io.StringIO(text), delimiter=delim))[: MAX_ROWS + 20]


def _is_utf8(data: bytes) -> bool:
    try:
        data.decode("utf-8-sig")
        return True
    except UnicodeDecodeError:
        return False


def _map_columns(header: list[Any]) -> dict[str, int]:
    found: dict[str, int] = {}
    names = [_norm(h) for h in header]
    for key, synonyms in HEADERS.items():
        for i, n in enumerate(names):
            if i in found.values() or not n:
                continue
            if n in synonyms or any(
                n.startswith(s + " ") or n.endswith(" " + s) for s in synonyms if len(s) > 3
            ):
                found[key] = i
                break
    return found


def _country(iban: str, country: str, text: str) -> str | None:
    for raw in (country,):
        if raw:
            raw = raw.strip()
            if len(raw) == 2 and raw.isalpha():
                return raw.upper()
            iso = nationality_iso(raw)
            if iso:
                return iso[0]
    for raw in (iban, text):
        m = IBAN.search((raw or "").replace(" ", "").upper())
        if m:
            return m.group(1)[:2]
    return None


def _getter(r: list[Any], mapping: dict[str, int]):
    def get(k: str) -> Any:
        return r[mapping[k]] if k in mapping and mapping[k] < len(r) else None

    return get


def parse(data: bytes, filename: str) -> tuple[list[dict[str, Any]], dict[str, str], list[str]]:
    rows = _rows(data, filename)
    warnings: list[str] = []
    head_idx, mapping = 0, {}
    for i, r in enumerate(rows[:20]):
        m = _map_columns(r)
        if "date" in m and ("amount" in m or "debit" in m or "credit" in m):
            head_idx, mapping = i, m
            break
    if not mapping:
        raise ValueError(
            "Columns not recognised: the file needs a date column and an amount column (or debit / credit columns)."
        )
    header = rows[head_idx]
    txs: list[dict[str, Any]] = []
    skipped = 0
    for r in rows[head_idx + 1 :]:
        if not any(str(c or "").strip() for c in r):
            continue
        get = _getter(r, mapping)
        d = _date(get("date"))
        amt = _amount(get("amount"))
        if amt is None and ("debit" in mapping or "credit" in mapping):
            deb, cre = _amount(get("debit")), _amount(get("credit"))
            amt = (abs(cre) if cre else 0.0) - (abs(deb) if deb else 0.0) if (deb or cre) else None
        if d is None or amt is None or amt == 0:
            skipped += 1
            continue
        desc = " ".join(str(x) for x in (get("description"), get("type")) if x)
        cp = str(get("counterparty") or "").strip()
        iban = str(get("iban") or "").strip()
        txs.append(
            {
                "idx": len(txs),
                "date": d,
                "amount": round(abs(amt), 2),
                "direction": "in" if amt > 0 else "out",
                "currency": str(get("currency") or "").strip().upper() or None,
                "counterparty": cp or None,
                "iban": iban or None,
                "country": _country(iban, str(get("country") or ""), f"{desc} {iban}"),
                "description": desc.strip()[:200] or None,
                "flags": [],
            }
        )
        if len(txs) >= MAX_ROWS:
            warnings.append(f"Only the first {MAX_ROWS} transactions were analysed.")
            break
    if skipped:
        warnings.append(f"{skipped} line(s) without a readable date or amount were skipped.")
    if not txs:
        raise ValueError("No transaction with a readable date and amount was found.")
    if "counterparty" not in mapping:
        warnings.append(
            "No counterparty column: counterparty screening and funnel checks are limited."
        )
    return txs, {k: str(header[i]) for k, i in mapping.items()}, warnings


def _alert(
    rule: str, severity: str, title: str, detail: str, rows: list[dict[str, Any]]
) -> dict[str, Any]:
    from app.legal import justification, refs_for

    for t in rows:
        if rule not in t["flags"]:
            t["flags"].append(rule)
    return {
        "legal": refs_for(rule),
        "justification_en": justification(rule, title, "en"),
        "justification_fr": justification(rule, title, "fr"),
        "rule": rule,
        "severity": severity,
        "title": title,
        "detail": detail,
        "count": len(rows),
        "amount": round(sum(t["amount"] for t in rows), 2),
        "rows": [t["idx"] for t in rows][:200],
    }


def _fmt(v: float) -> str:
    return f"{v:,.0f}".replace(",", "'")


def analyse(
    data: bytes,
    filename: str,
    profile: dict[str, Any] | None = None,
    screen: Any = None,
) -> dict[str, Any]:
    """`profile`: {country, volume, cash, threshold} from the KYC file. `screen(entities)`
    screens the counterparties in one go (the enabled sanctions / watchlist connectors) and
    returns (hits, notes); None skips it."""
    if len(data) > MAX_BYTES:
        raise ValueError("File too large (5 MB maximum).")
    profile = profile or {}
    txs, mapping, warnings = parse(data, filename)
    jur = get_jurisdictions()
    risk = get_country_risk()
    t = get_risk_config().thresholds
    home = (profile.get("country") or "").upper()
    ccys = Counter(x["currency"] for x in txs if x["currency"])
    ccy = ccys.most_common(1)[0][0] if ccys else None
    # CHF 15 000 (Swiss identification threshold for cash); 10 000 in EUR / USD / others
    swiss = ccy == "CHF" or (ccy is None and home == "CH")
    threshold = float(profile.get("threshold") or (15_000 if swiss else 10_000))
    alerts: list[dict[str, Any]] = []
    txs.sort(key=lambda x: (x["date"], x["idx"]))

    # 1. Structuring: just below the threshold, or split over a few days
    near = [x for x in txs if 0.9 * threshold <= x["amount"] < threshold]
    if len(near) >= 2:
        alerts.append(
            _alert(
                "structuring",
                "high",
                f"{len(near)} amounts just below the {_fmt(threshold)} threshold",
                f"Transactions between {_fmt(0.9 * threshold)} and {_fmt(threshold)}: a classic way to avoid identification or reporting.",
                near,
            )
        )
    by_cp_week: dict[tuple, list[dict[str, Any]]] = defaultdict(list)
    for x in txs:
        if x["amount"] < threshold and x["counterparty"]:
            by_cp_week[
                (x["counterparty"].lower(), x["direction"], x["date"].isocalendar()[:2])
            ].append(x)
    split = [
        g for g in by_cp_week.values() if len(g) >= 3 and sum(y["amount"] for y in g) >= threshold
    ]
    if split:
        rows = [y for g in split for y in g]
        alerts.append(
            _alert(
                "split_payments",
                "high",
                f"{len(split)} group(s) of payments split within a week",
                "Several payments to or from the same counterparty in the same week, each below the threshold but above it together.",
                rows,
            )
        )

    # 2. Cash
    cash = [x for x in txs if x["description"] and CASH.search(unidecode(x["description"]))]
    if cash:
        total = sum(x["amount"] for x in cash)
        share = total / max(1.0, sum(x["amount"] for x in txs))
        declared_none = profile.get("cash") == "none"
        alerts.append(
            _alert(
                "cash",
                "high" if declared_none or share >= 0.2 else "medium",
                f"Cash: {len(cash)} operations, {_fmt(total)} ({share:.0%} of the volume)",
                ("The KYC profile declares no use of cash. " if declared_none else "")
                + "Ask for the origin of cash deposits and the use of cash withdrawals.",
                cash,
            )
        )

    # 3. Round amounts
    big = [x for x in txs if x["amount"] >= 5000]
    rounds = [x for x in big if x["amount"] % 1000 == 0]
    if len(big) >= 5 and len(rounds) / len(big) >= 0.3:
        alerts.append(
            _alert(
                "round_amounts",
                "medium",
                f"{len(rounds)} of {len(big)} large transfers are round amounts",
                "Round amounts are unusual for commercial payments (invoices rarely end in 000): check the underlying contracts.",
                rounds,
            )
        )

    # 4. Pass-through: out within 3 days of an in, for 90-100 % of it
    ins = [x for x in txs if x["direction"] == "in" and x["amount"] >= 5000]
    outs = [x for x in txs if x["direction"] == "out"]
    used: set[int] = set()
    pairs: list[dict[str, Any]] = []
    for i in ins:
        for o in outs:
            if o["idx"] in used or not (i["date"] <= o["date"] <= i["date"] + timedelta(days=3)):
                continue
            if 0.9 * i["amount"] <= o["amount"] <= 1.0 * i["amount"]:
                used.add(o["idx"])
                pairs += [i, o]
                break
    inflow = sum(x["amount"] for x in txs if x["direction"] == "in")
    outflow = sum(x["amount"] for x in txs if x["direction"] == "out")
    passed = sum(x["amount"] for x in pairs if x["direction"] == "in")
    if pairs and passed >= 0.25 * max(inflow, 1):
        alerts.append(
            _alert(
                "pass_through",
                "high",
                f"Pass-through: {_fmt(passed)} left within 3 days of arriving ({passed / max(inflow, 1):.0%} of inflows)",
                "Funds credited then debited for nearly the same amount within days: the account may be used as a transit (layering).",
                pairs,
            )
        )

    # 5. Geography
    countries: dict[str, dict[str, Any]] = {}
    for x in txs:
        c = x["country"]
        if not c:
            continue
        row = countries.setdefault(
            c, {"iso": c, "name": jur.name(c), "amount": 0.0, "count": 0, "risk": []}
        )
        row["amount"] += x["amount"]
        row["count"] += 1
    geo_alerts: dict[str, tuple[str, str]] = {}
    for c, row in countries.items():
        reasons = []
        sev = None
        if c in jur.fatf_blacklist:
            reasons.append("FATF call for action")
            sev = "critical"
        if c in BROADLY_SANCTIONED:
            reasons.append("broad sanctions programmes")
            sev = sev or "high"
        if c in jur.fatf_greylist:
            reasons.append("FATF increased monitoring")
            sev = sev or "medium"
        if c in jur.eu_tax_blacklist:
            reasons.append("EU tax list")
            sev = sev or "medium"
        if c in jur.offshore_centres:
            reasons.append("offshore centre")
            sev = sev or "medium"
        basel = risk.get(c).get("basel_aml_score")
        if basel is not None and basel >= t.get("basel_high_score", 6.0):
            reasons.append(f"Basel AML Index {basel}")
            sev = sev or "medium"
        row["risk"] = reasons
        if sev:
            geo_alerts[c] = (sev, ", ".join(reasons))
    order = {"critical": 0, "high": 1, "medium": 2, "low": 3}
    for c, (sev, why) in sorted(geo_alerts.items(), key=lambda kv: order[kv[1][0]]):
        rows = [x for x in txs if x["country"] == c]
        alerts.append(
            _alert(
                "high_risk_geography",
                sev,
                f"{countries[c]['name']}: {len(rows)} transaction(s), {_fmt(countries[c]['amount'])}",
                f"Counterparties in {countries[c]['name']} ({why}). Document the business reason and the goods or services involved.",
                rows,
            )
        )

    # 6. Crypto platforms
    crypto = [
        x for x in txs if CRYPTO.search(f"{x['counterparty'] or ''} {x['description'] or ''}")
    ]
    if crypto:
        alerts.append(
            _alert(
                "crypto",
                "medium",
                f"Crypto-asset platforms: {len(crypto)} transfer(s), {_fmt(sum(x['amount'] for x in crypto))}",
                "Transfers with crypto exchanges: ask for the wallet addresses and the origin of the crypto-assets, and screen the wallets.",
                crypto,
            )
        )

    # 7. Monthly spikes
    monthly: dict[str, dict[str, float]] = defaultdict(lambda: {"in": 0.0, "out": 0.0})
    for x in txs:
        monthly[x["date"].strftime("%Y-%m")][x["direction"]] += x["amount"]
    months = sorted(monthly)
    totals = {m: monthly[m]["in"] + monthly[m]["out"] for m in months}
    if len(months) >= 3:
        for m in months:
            others = [v for k, v in totals.items() if k != m]
            med = statistics.median(others)
            if med > 0 and totals[m] > 3 * med:
                rows = [x for x in txs if x["date"].strftime("%Y-%m") == m]
                alerts.append(
                    _alert(
                        "activity_spike",
                        "medium",
                        f"Activity spike in {m}: {_fmt(totals[m])} (×{totals[m] / med:.1f} the usual month)",
                        "A sudden rise in volume: ask what explains it (contract, sale of an asset, new business).",
                        rows,
                    )
                )

    # 8. Declared profile
    first, last = txs[0]["date"], txs[-1]["date"]
    days = max(30, (last - first).days + 1)
    annual = inflow * 365 / days
    band = VOLUME_BANDS.get(profile.get("volume") or "")
    if band and band[1] and annual > 1.5 * band[1]:
        alerts.append(
            _alert(
                "profile_mismatch",
                "high",
                f"Inflows of {_fmt(annual)} a year, above the declared profile (up to {_fmt(band[1])})",
                "The activity exceeds what the client declared at onboarding: update the profile and ask for the reason.",
                [x for x in txs if x["direction"] == "in"],
            )
        )

    # 9. Funnel: many senders, few beneficiaries
    senders = {
        x["counterparty"].lower() for x in txs if x["direction"] == "in" and x["counterparty"]
    }
    receivers = Counter(
        x["counterparty"].lower() for x in txs if x["direction"] == "out" and x["counterparty"]
    )
    if len(senders) >= 10 and receivers and len(receivers) <= 2 and outflow >= 0.7 * inflow:
        alerts.append(
            _alert(
                "funnel",
                "high",
                f"Funnel account: {len(senders)} different senders, outflows to {len(receivers)} beneficiar{'y' if len(receivers) == 1 else 'ies'}",
                "Many incoming payments from different people forwarded to very few beneficiaries: a money-mule or collection-account pattern.",
                [x for x in txs if x["counterparty"]],
            )
        )

    # 10. Counterparty screening
    by_cp: dict[str, dict[str, Any]] = {}
    for x in txs:
        if not x["counterparty"]:
            continue
        row = by_cp.setdefault(
            x["counterparty"].lower(),
            {
                "name": x["counterparty"],
                "in": 0.0,
                "out": 0.0,
                "count": 0,
                "country": x["country"],
                "hits": [],
            },
        )
        row[x["direction"]] += x["amount"]
        row["count"] += 1
    top = sorted(by_cp.values(), key=lambda r: -(r["in"] + r["out"]))
    screened = 0
    if screen is not None:
        entities = [
            Entity(
                id=f"stmt:{i}",
                type=EntityType.COMPANY if LEGAL_WORDS.search(row["name"]) else EntityType.PERSON,
                name=row["name"],
                jurisdiction=row["country"],
            )
            for i, row in enumerate(top[:MAX_SCREENED])
        ]
        try:
            hits, notes = screen(entities)
        except Exception as exc:  # noqa: BLE001 - a source outage must not stop the analysis
            hits, notes = [], [f"Screening unavailable: {str(exc)[:80]}"]
        warnings += notes
        screened = len(entities)
        by_entity: dict[str, list] = defaultdict(list)
        for h in hits:
            by_entity[h.entity_id].append(h)
        for ent, row in zip(entities, top, strict=False):
            row["hits"] = [
                {
                    "dataset": h.dataset,
                    "matched_name": h.matched_name,
                    "score": h.score,
                    "list_type": h.list_type.value,
                }
                for h in sorted(by_entity[ent.id], key=lambda h: -h.score)
                if h.score >= t["possible_match_score"]
                and h.triage not in ("namesake", "dismissed")
            ][:5]
        flagged = [r for r in top if r["hits"]]
        for r in flagged:
            best = max(r["hits"], key=lambda h: h["score"])
            rows = [x for x in txs if (x["counterparty"] or "").lower() == r["name"].lower()]
            strong = best["score"] >= t["strong_match_score"] and best["list_type"] == "sanction"
            alerts.append(
                _alert(
                    "counterparty_screening",
                    "critical" if strong else "high",
                    f"{r['name']}: possible match on {best['dataset']} ({best['score']:g})",
                    f"Listed as {best['matched_name']}. Confirm or rule out the match before executing any further payment.",
                    rows,
                )
            )
    alerts.sort(key=lambda a: (order.get(a["severity"], 4), -a["amount"]))
    currencies = Counter(x["currency"] for x in txs if x["currency"])
    return {
        "filename": filename,
        "mapping": mapping,
        "warnings": warnings,
        "threshold": threshold,
        "summary": {
            "transactions": len(txs),
            "period_from": first.isoformat(),
            "period_to": last.isoformat(),
            "currency": currencies.most_common(1)[0][0] if currencies else None,
            "currencies": dict(currencies),
            "inflow": round(inflow, 2),
            "outflow": round(outflow, 2),
            "count_in": sum(x["direction"] == "in" for x in txs),
            "count_out": sum(x["direction"] == "out" for x in txs),
            "annualised_inflow": round(annual, 2),
            "counterparties": len(by_cp),
            "screened": screened,
            "flagged_transactions": sum(bool(x["flags"]) for x in txs),
        },
        "monthly": [
            {"month": m, "in": round(monthly[m]["in"], 2), "out": round(monthly[m]["out"], 2)}
            for m in months
        ],
        "countries": sorted(
            ({**c, "amount": round(c["amount"], 2)} for c in countries.values()),
            key=lambda c: -c["amount"],
        ),
        "counterparties": [
            {**r, "in": round(r["in"], 2), "out": round(r["out"], 2)} for r in top[:50]
        ],
        "alerts": alerts,
        "transactions": [{**x, "date": x["date"].isoformat()} for x in txs[:3000]],
    }
